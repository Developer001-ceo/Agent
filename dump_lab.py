"""Dump Lab7_8-v1.xlsx structure so we can find Task 1 instructions."""
import openpyxl
import sys
import json
import os

PATH = r"C:\Users\prasa\OneDrive\Pictures\Lab 7 & 8\Lab 7 & 8\Lab7_8-v1.xlsx"

print(f"[info] file size = {os.path.getsize(PATH):,} bytes")
print(f"[info] opening read-only ...")
wb = openpyxl.load_workbook(PATH, data_only=True, read_only=True)
print(f"[info] sheets: {wb.sheetnames}")
print("=" * 80)

MAX_ROWS = 120
MAX_COLS = 16

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n### SHEET: {name!r}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print("-" * 80)
    row_iter = ws.iter_rows(min_row=1, max_row=min(MAX_ROWS, ws.max_row or 1),
                            max_col=min(MAX_COLS, ws.max_column or 1),
                            values_only=True)
    for i, row in enumerate(row_iter, 1):
        # compress Nones to keep output readable
        cells = [("" if v is None else str(v)) for v in row]
        if any(c.strip() for c in cells):
            print(f"R{i:>3}: {cells}")
    # Also dump any images / drawings count if available
    try:
        n_img = len(getattr(ws, "_images", []) or [])
    except Exception:
        n_img = 0
    print(f"[sheet {name}] images={n_img}")

print("\n[DONE]")
